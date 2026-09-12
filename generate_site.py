#!/usr/bin/env python3
import os
import re
import json
import glob
import openpyxl

def extract_youtube_id(url):
    if not url or not str(url).strip():
        return None
    url = str(url).strip()
    match = re.search(r'(?:v=|\/|embed\/|youtu\.be\/)([0-9A-Za-z_-]{11})', url)
    return match.group(1) if match else None

def generate_html(excel_path, output_html_path):
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb["Course_Content"] if "Course_Content" in wb.sheetnames else wb.active
    
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        print("Error: Excel sheet is empty!")
        return

    headers = [str(h).strip() if h else "" for h in rows[0]]
    data_rows = rows[1:]
    
    units = {}
    for idx, r in enumerate(data_rows, start=1):
        if not any(r):
            continue
        row_dict = {headers[i]: (r[i] if i < len(r) and r[i] is not None else "") for i in range(len(headers))}
        unit_name = str(row_dict.get("Unit", "General Unit")).strip()
        if not unit_name:
            unit_name = "General Unit"
        if unit_name not in units:
            units[unit_name] = []
        row_dict["index"] = idx
        row_dict["yt_id"] = extract_youtube_id(row_dict.get("YouTube_Link", ""))
        units[unit_name].append(row_dict)

    units_json = json.dumps(units)

    html_template = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Interactive Mathematics Course Hub</title>
  
  <!-- Tailwind CSS CDN -->
  <script src="https://cdn.tailwindcss.com"></script>
  
  <!-- MathJax for LaTeX Rendering -->
  <script>
    window.MathJax = {{
      tex: {{
        inlineMath: [['$', '$'], ['\\\\(', '\\\\)']],
        displayMath: [['$$', '$$'], ['\\\\[', '\\\\]']]
      }},
      svg: {{ fontCache: 'global' }}
    }};
  </script>
  <script id="MathJax-script" async src="https://cdn.jsdelivr.net/npm/mathjax@3/es5/tex-mml-chtml.js"></script>
  
  <!-- JSXGraph CSS and JS -->
  <link rel="stylesheet" type="text/css" href="https://cdn.jsdelivr.net/npm/jsxgraph/distrib/jsxgraph.css" />
  <script type="text/javascript" src="https://cdn.jsdelivr.net/npm/jsxgraph/distrib/jsxgraphcore.js"></script>

  <!-- Lucide Icons -->
  <script src="https://unpkg.com/lucide@latest"></script>

  <style>
    .jxgbox {{
      width: 100%;
      height: 380px;
      border-radius: 0.75rem;
      border: 1px solid #e2e8f0;
      box-shadow: 0 1px 3px 0 rgba(0, 0, 0, 0.05);
      background-color: #ffffff;
    }}
  </style>
</head>
<body class="bg-slate-50 text-slate-800 antialiased min-h-screen flex flex-col font-sans">

  <!-- Header -->
  <header class="bg-indigo-900 text-white shadow-md sticky top-0 z-50">
    <div class="max-w-7xl mx-auto px-4 py-4 flex justify-between items-center">
      <div class="flex items-center space-x-3">
        <div class="bg-indigo-700 p-2 rounded-lg">
          <i data-lucide="book-open" class="w-6 h-6 text-indigo-200"></i>
        </div>
        <div>
          <h1 class="text-xl font-bold tracking-tight">MathEd 548 Course Hub</h1>
          <p class="text-xs text-indigo-200">Interactive Mathematics &amp; Dynamic Geometry</p>
        </div>
      </div>
      <div class="flex items-center space-x-2">
        <a href="https://github.com/bedprasad/MathEd548_01/releases/download/latest/main.pdf" target="_blank" class="inline-flex items-center px-3 py-1.5 bg-indigo-700 hover:bg-indigo-600 text-xs font-semibold rounded-md shadow-sm transition">
          <i data-lucide="file-down" class="w-4 h-4 mr-1.5"></i> Course PDF
        </a>
      </div>
    </div>
  </header>

  <!-- Main Content Layout -->
  <div class="max-w-7xl mx-auto px-4 py-6 flex-1 w-full grid grid-cols-1 lg:grid-cols-4 gap-6">
    
    <!-- Sidebar / Unit Navigation -->
    <aside class="lg:col-span-1 bg-white rounded-xl shadow-sm border border-slate-200 p-4 h-fit sticky top-24">
      <h2 class="text-xs font-bold text-slate-400 uppercase tracking-wider mb-3">Course Curriculum</h2>
      <div id="unit-nav" class="space-y-4">
        <!-- Injected via JS -->
      </div>
    </aside>

    <!-- Main Lesson Area -->
    <main class="lg:col-span-3 space-y-6">
      <div id="lesson-container">
        <!-- Active Lesson Injected via JS -->
      </div>
    </main>
  </div>

  <!-- Footer -->
  <footer class="bg-white border-t border-slate-200 py-4 mt-12 text-center text-xs text-slate-500">
    Interactive Course Module &bull; Powered by JSXGraph, MathJax, and GitHub Pages
  </footer>

  <script>
    const courseData = {units_json};
    let activeUnit = Object.keys(courseData)[0];
    let activeTopicIdx = 0;

    function renderNav() {{
      const nav = document.getElementById('unit-nav');
      nav.innerHTML = '';
      
      Object.keys(courseData).forEach((unitKey, uIdx) => {{
        const unitDiv = document.createElement('div');
        unitDiv.className = 'border-b border-slate-100 pb-3 last:border-0 last:pb-0';
        
        const title = document.createElement('div');
        title.className = 'font-semibold text-sm text-slate-700 mb-2 flex items-center justify-between cursor-pointer';
        title.innerHTML = `<span>${{unitKey}}</span>`;
        unitDiv.appendChild(title);

        const list = document.createElement('ul');
        list.className = 'space-y-1 pl-2';

        courseData[unitKey].forEach((topic, tIdx) => {{
          const li = document.createElement('li');
          const isSelected = (unitKey === activeUnit && tIdx === activeTopicIdx);
          li.className = `text-xs px-2.5 py-1.5 rounded-lg cursor-pointer transition flex items-center justify-between ${{
            isSelected ? 'bg-indigo-50 text-indigo-700 font-semibold border-l-4 border-indigo-600' : 'text-slate-600 hover:bg-slate-50'
          }}`;
          li.innerHTML = `<span>${{topic.Topic_Name}}</span>`;
          li.onclick = () => {{
            activeUnit = unitKey;
            activeTopicIdx = tIdx;
            renderNav();
            renderLesson();
          }};
          list.appendChild(li);
        }});

        unitDiv.appendChild(list);
        nav.appendChild(unitDiv);
      }});
      lucide.createIcons();
    }}

    function renderLesson() {{
      const container = document.getElementById('lesson-container');
      const topic = courseData[activeUnit][activeTopicIdx];
      
      const boxId = `jxgbox_${{topic.index}}`;
      
      let quizOptionsHtml = '';
      if (topic.Quiz_Options) {{
        const opts = topic.Quiz_Options.split('|').map(o => o.trim());
        quizOptionsHtml = opts.map((opt, i) => `
          <label class="flex items-center space-x-3 p-2.5 bg-slate-50 rounded-lg border border-slate-200 cursor-pointer hover:bg-indigo-50 transition">
            <input type="radio" name="quiz_opt_${{topic.index}}" value="${{opt.replace(/"/g, '&quot;')}}" class="text-indigo-600 focus:ring-indigo-500">
            <span class="text-sm text-slate-700">${{opt}}</span>
          </label>
        `).join('');
      }}

      container.innerHTML = `
        <div class="bg-white rounded-xl shadow-sm border border-slate-200 p-6 space-y-6">
          
          <!-- Topic Header -->
          <div class="border-b border-slate-100 pb-4">
            <div class="text-xs font-semibold text-indigo-600 tracking-wide uppercase mb-1">${{activeUnit}}</div>
            <h2 class="text-2xl font-bold text-slate-800">${{topic.Topic_Name}}</h2>
          </div>

          <!-- HTML Content -->
          ${{topic.HTML_Content ? `
            <div class="prose prose-slate max-w-none text-slate-700 leading-relaxed text-sm">
              ${{topic.HTML_Content}}
            </div>
          ` : ''}}

          <!-- LaTeX Math Section -->
          ${{topic.LaTeX_Math ? `
            <div class="bg-indigo-50/50 border border-indigo-100 rounded-xl p-4 text-slate-800">
              <div class="text-xs font-bold text-indigo-800 uppercase tracking-wider mb-2 flex items-center">
                <i data-lucide="function-square" class="w-4 h-4 mr-1.5"></i> Mathematical Formulation
              </div>
              <div class="math-content whitespace-pre-wrap text-sm leading-loose">
                ${{topic.LaTeX_Math}}
              </div>
            </div>
          ` : ''}}

          <!-- Interactive JSXGraph Simulation -->
          ${{topic.JSXGraph_Code ? `
            <div class="space-y-2">
              <div class="flex items-center justify-between">
                <span class="text-xs font-bold text-slate-500 uppercase tracking-wider flex items-center">
                  <i data-lucide="sparkles" class="w-4 h-4 mr-1.5 text-amber-500"></i> Interactive Graph Simulation
                </span>
                <span class="text-xs text-slate-400">Drag points &amp; sliders to explore</span>
              </div>
              <div id="${{boxId}}" class="jxgbox"></div>
            </div>
          ` : ''}}

          <!-- YouTube Video Embed -->
          ${{topic.yt_id ? `
            <div class="space-y-2">
              <div class="text-xs font-bold text-slate-500 uppercase tracking-wider flex items-center">
                <i data-lucide="youtube" class="w-4 h-4 mr-1.5 text-red-500"></i> Video Lecture
              </div>
              <div class="aspect-video w-full rounded-xl overflow-hidden shadow-inner bg-slate-900">
                <iframe class="w-full h-full" src="https://www.youtube.com/embed/${{topic.yt_id}}" frameborder="0" allowfullscreen></iframe>
              </div>
            </div>
          ` : ''}}

          <!-- PDF Reference Link -->
          ${{topic.PDF_Link ? `
            <div class="flex items-center justify-between p-3.5 bg-slate-50 rounded-xl border border-slate-200">
              <div class="flex items-center space-x-3">
                <div class="p-2 bg-red-100 text-red-600 rounded-lg">
                  <i data-lucide="file-text" class="w-5 h-5"></i>
                </div>
                <div>
                  <div class="text-xs font-bold text-slate-700">Course Notes &amp; Derivations</div>
                  <div class="text-xs text-slate-500">Download the latest compiled PDF for this module</div>
                </div>
              </div>
              <a href="${{topic.PDF_Link}}" target="_blank" class="px-3.5 py-1.5 bg-white hover:bg-slate-100 text-slate-700 border border-slate-300 font-semibold text-xs rounded-lg shadow-sm transition">
                Open PDF
              </a>
            </div>
          ` : ''}}

          <!-- Self Assessment Question -->
          ${{topic.SA_Question ? `
            <div class="border border-slate-200 rounded-xl p-5 bg-gradient-to-br from-slate-50 to-white">
              <div class="text-xs font-bold text-indigo-600 uppercase tracking-wider mb-2 flex items-center">
                <i data-lucide="help-circle" class="w-4 h-4 mr-1.5"></i> Self-Assessment Prompt
              </div>
              <p class="text-sm font-medium text-slate-800 mb-3">${{topic.SA_Question}}</p>
              <button onclick="toggleAnswer('sa_ans_${{topic.index}}')" class="text-xs font-semibold text-indigo-600 hover:text-indigo-800 transition flex items-center">
                <i data-lucide="eye" class="w-3.5 h-3.5 mr-1"></i> Toggle Solution
              </button>
              <div id="sa_ans_${{topic.index}}" class="hidden mt-3 p-3 bg-indigo-50/70 border border-indigo-100 rounded-lg text-xs text-slate-700 leading-relaxed">
                ${{topic.SA_Answer}}
              </div>
            </div>
          ` : ''}}

          <!-- Interactive Quiz -->
          ${{topic.Quiz_Question ? `
            <div class="border border-indigo-100 bg-indigo-50/30 rounded-xl p-5 space-y-4">
              <div class="flex justify-between items-center">
                <span class="text-xs font-bold text-indigo-700 uppercase tracking-wider flex items-center">
                  <i data-lucide="check-circle-2" class="w-4 h-4 mr-1.5"></i> Knowledge Check Quiz
                </span>
              </div>
              <p class="text-sm font-semibold text-slate-800">${{topic.Quiz_Question}}</p>
              <div class="space-y-2">
                ${{quizOptionsHtml}}
              </div>
              <button onclick="checkQuizAnswer(${{topic.index}}, '${{encodeURIComponent(topic.Quiz_Correct_Answer)}}', '${{encodeURIComponent(topic.Quiz_Explanation)}}')" class="px-4 py-2 bg-indigo-600 hover:bg-indigo-700 text-white font-semibold text-xs rounded-lg shadow-sm transition">
                Check Answer
              </button>
              <div id="quiz_feedback_${{topic.index}}" class="hidden text-xs p-3 rounded-lg"></div>
            </div>
          ` : ''}}

        </div>
      `;

      lucide.createIcons();
      if (window.MathJax) {{
        MathJax.typesetPromise([container]);
      }}

      // Execute JSXGraph Script if present
      if (topic.JSXGraph_Code) {{
        try {{
          const scriptFn = new Function(topic.JSXGraph_Code);
          setTimeout(() => {{
            scriptFn();
          }}, 50);
        }} catch (e) {{
          console.error("JSXGraph Error:", e);
        }}
      }}
    }}

    function toggleAnswer(id) {{
      const el = document.getElementById(id);
      el.classList.toggle('hidden');
    }}

    function checkQuizAnswer(index, encCorrect, encExp) {{
      const correct = decodeURIComponent(encCorrect).trim();
      const explanation = decodeURIComponent(encExp);
      const selected = document.querySelector(`input[name="quiz_opt_${{index}}"]:checked`);
      const feedback = document.getElementById(`quiz_feedback_${{index}}`);
      
      feedback.classList.remove('hidden', 'bg-emerald-50', 'text-emerald-800', 'border-emerald-200', 'bg-red-50', 'text-red-800', 'border-red-200');
      
      if (!selected) {{
        feedback.className = 'text-xs p-3 rounded-lg border bg-amber-50 text-amber-800 border-amber-200';
        feedback.innerHTML = '<strong>Please select an option first.</strong>';
        return;
      }}
      
      if (selected.value.trim() === correct) {{
        feedback.className = 'text-xs p-3 rounded-lg border bg-emerald-50 text-emerald-800 border-emerald-200';
        feedback.innerHTML = `<strong>✔ Correct!</strong> ${{explanation}}`;
      }} else {{
        feedback.className = 'text-xs p-3 rounded-lg border bg-red-50 text-red-800 border-red-200';
        feedback.innerHTML = `<strong>✖ Not quite.</strong> The correct answer is <em>${{correct}}</em>.<br>${{explanation}}`;
      }}
      
      if (window.MathJax) {{
        MathJax.typesetPromise([feedback]);
      }}
    }}

    // Init
    window.addEventListener('DOMContentLoaded', () => {{
      renderNav();
      renderLesson();
    }});
  </script>
</body>
</html>
"""
    os.makedirs(os.path.dirname(output_html_path), exist_ok=True)
    with open(output_html_path, "w", encoding="utf-8") as f:
        f.write(html_template)
    print(f"Successfully generated {output_html_path}")

if __name__ == "__main__":
    excel_files = glob.glob("*.xlsx")
    target_excel = "course_sample.xlsx"
    if "course_data.xlsx" in excel_files:
        target_excel = "course_data.xlsx"
    elif excel_files:
        target_excel = excel_files[0]
    
    print(f"Reading from: {target_excel}")
    generate_html(target_excel, "docs/index.html")
